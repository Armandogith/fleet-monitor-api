from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from pydantic import BaseModel
from datetime import date, timedelta
from typing import Optional
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from app.core.database import AsyncSessionLocal
from app.models.schema import Driver, Document, Company, NotificationLog
import openpyxl
import io

router = APIRouter()


# ── Schemas ───────────────────────────────────────────────────────────────────

class DriverCreate(BaseModel):
    company_id: int
    name: str
    phone_number: str


class DocumentCreate(BaseModel):
    driver_id: int
    doc_type: str
    expiration_date: date


class CompanyCreate(BaseModel):
    name: str
    responsible_email: Optional[str] = None
    responsible_whatsapp: Optional[str] = None


class CompanyUpdate(BaseModel):
    responsible_email: Optional[str] = None
    responsible_whatsapp: Optional[str] = None


# ── DB dependency ─────────────────────────────────────────────────────────────

async def get_db():
    async with AsyncSessionLocal() as session:
        yield session


# ── EMPRESAS ──────────────────────────────────────────────────────────────────

@router.get("/companies")
async def list_companies(db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Company))
    companies = result.scalars().all()
    return [
        {
            "id": c.id,
            "name": c.name,
            "responsible_email": c.responsible_email,
            "plan_type": c.plan_type or "demo",
            "plan_expires_at": str(c.plan_expires_at) if c.plan_expires_at else None,
        }
        for c in companies
    ]


@router.post("/companies")
async def create_company(company: CompanyCreate, db: AsyncSession = Depends(get_db)):
    new_company = Company(
        name=company.name,
        responsible_email=company.responsible_email,
        responsible_whatsapp=company.responsible_whatsapp,
        plan_type="demo",
        plan_expires_at=date.today() + timedelta(days=7),
    )
    db.add(new_company)
    await db.commit()
    await db.refresh(new_company)
    return {"status": "success", "company_id": new_company.id}


@router.put("/companies/{company_id}/responsible")
async def update_responsible(company_id: int, data: CompanyUpdate, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Company).where(Company.id == company_id))
    company = result.scalar_one_or_none()
    if not company:
        raise HTTPException(status_code=404, detail="Empresa não encontrada")
    if data.responsible_email is not None:
        company.responsible_email = data.responsible_email
    if data.responsible_whatsapp is not None:
        company.responsible_whatsapp = data.responsible_whatsapp
    await db.commit()
    return {"status": "success"}


# ── MOTORISTAS ────────────────────────────────────────────────────────────────

@router.get("/drivers")
async def list_drivers(db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Driver))
    drivers = result.scalars().all()
    return [
        {"id": d.id, "name": d.name, "phone_number": d.phone_number, "company_id": d.company_id}
        for d in drivers
    ]


@router.post("/drivers")
async def create_driver(driver: DriverCreate, db: AsyncSession = Depends(get_db)):
    new_driver = Driver(
        company_id=driver.company_id,
        name=driver.name,
        phone_number=driver.phone_number,
    )
    db.add(new_driver)
    await db.commit()
    await db.refresh(new_driver)
    return {"status": "success", "driver_id": new_driver.id}


# ── DOCUMENTOS ────────────────────────────────────────────────────────────────

@router.get("/documents")
async def list_documents(db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Document))
    docs = result.scalars().all()
    return [
        {
            "id": d.id,
            "driver_id": d.driver_id,
            "doc_type": d.doc_type,
            "expiration_date": str(d.expiration_date),
        }
        for d in docs
    ]


@router.post("/documents")
async def create_document(doc: DocumentCreate, db: AsyncSession = Depends(get_db)):
    new_doc = Document(
        driver_id=doc.driver_id,
        doc_type=doc.doc_type,
        expiration_date=doc.expiration_date,
    )
    db.add(new_doc)
    await db.commit()
    await db.refresh(new_doc)
    return {"status": "success", "document_id": new_doc.id}


# ── UPLOAD DE PLANILHA ────────────────────────────────────────────────────────

@router.post("/upload-drivers")
async def upload_drivers(
    company_id: int,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Envie um arquivo .xlsx ou .xls")

    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents))
    ws = wb.active
    criados = 0
    erros = []

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        try:
            nome, telefone, tipo_doc, data_venc = row[0], row[1], row[2], row[3]
            if not all([nome, telefone, tipo_doc, data_venc]):
                erros.append(f"Linha {i}: dados incompletos")
                continue
            if isinstance(data_venc, str):
                from datetime import datetime
                data_venc = datetime.strptime(data_venc.strip(), "%d/%m/%Y").date()
            driver = Driver(
                company_id=company_id,
                name=str(nome).strip(),
                phone_number=str(telefone).strip(),
            )
            db.add(driver)
            await db.flush()
            doc = Document(
                driver_id=driver.id,
                doc_type=str(tipo_doc).strip(),
                expiration_date=data_venc,
            )
            db.add(doc)
            criados += 1
        except Exception as e:
            erros.append(f"Linha {i}: {str(e)}")

    await db.commit()
    return {"status": "success", "motoristas_criados": criados, "erros": erros}


# ── HISTÓRICO ─────────────────────────────────────────────────────────────────

async def _list_logs(db: AsyncSession):
    result = await db.execute(
        select(NotificationLog).order_by(NotificationLog.timestamp.desc())
    )
    logs = result.scalars().all()
    return [
        {
            "id": l.id,
            "document_id": l.document_id,
            "status": l.status,
            "error_message": l.error_message,
            "timestamp": str(l.timestamp),
        }
        for l in logs
    ]


@router.get("/notification-logs")
async def list_logs(db: AsyncSession = Depends(get_db)):
    return await _list_logs(db)


@router.get("/notifications")
async def list_notifications(db: AsyncSession = Depends(get_db)):
    return await _list_logs(db)


# ── DISPARO DO JOB ────────────────────────────────────────────────────────────

async def _run_job():
    from app.graph.workflow import create_workflow
    workflow = create_workflow()
    result = await workflow.ainvoke({
        "batch_id": "",
        "pending_tasks": [],
        "send_results": [],
    })
    return result.get("send_results", [])


@router.post("/trigger-job")
async def trigger_job():
    send_results = await _run_job()
    return {"status": "ok", "send_results": send_results}


@router.post("/trigger-scheduler")
async def trigger_scheduler():
    send_results = await _run_job()
    return {"message": f"Job executado! {len(send_results)} notificações enviadas."}
@router.get("/download-template")
async def download_template():
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    import io
    wb = Workbook()
    ws = wb.active
    ws.title = "Modelo Importacao"
    ws.append(["Nome Completo", "Telefone (Apenas Numeros)", "Tipo de Documento", "Data de Vencimento (DD/MM/AAAA)"])
    ws.append(["João da Silva", "11999990000", "CNH", "15/10/2026"])
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 35
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=modelo_importacao.xlsx"}
    )
